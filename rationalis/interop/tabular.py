"""Esportazione dei movimenti (Excel, CSV, HTML) e importazione da fogli Excel.

Formato delle colonne (uguale in uscita e in ingresso):
    data, gruppo, sottogruppo, risorsa, valuta, entrata, uscita, s, modalità, destinazione,
    soggetto, descrizione, rif, c
Un giroconto è scritto come più righe consecutive: la prima (il prelievo) ha gruppo «<giroconto>»,
le successive (i versamenti) hanno il gruppo vuoto.

Regole di importazione da Excel:
- l'intestazione deve trovarsi nella prima colonna, entro le prime 20 righe;
- colonne in qualsiasi ordine, non tutte necessarie; obbligatorie data, gruppo, risorsa;
- l'importazione termina alla prima riga vuota;
- la colonna valuta è ignorata; le voci di tabella mancanti vengono create;
- i movimenti si aggiungono sempre a quelli esistenti.
"""
import csv
import re
import html
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from .. import movements as mv
from .. import transfers as tr
from ..db import Archive, ValidationError
from ..money import ZERO, fmt_amount, parse_amount
from . import ImportReport
from .resolver import TRANSFER_MARK, Resolver

COLUMNS = ["data", "gruppo", "sottogruppo", "risorsa", "valuta", "entrata", "uscita", "s", "modalità",
           "destinazione", "soggetto", "descrizione", "rif", "c"]
ALIASES = {"modalita": "modalità", "saldato": "s", "controllo": "c", "riferimento": "rif", "giorno": "data"}
HEADER_SEARCH_ROWS = 20


# --- esportazione -----------------------------------------------------------
def export_rows(arc: Archive, start: date, end: date) -> list[list]:
    """Righe da esportare (valori Python), giroconti compresi secondo la convenzione del modulo."""
    out = []
    name = arc.name_of
    items = []
    for m in mv.list_period(arc, start, end, include_transfers=False):
        items.append((m.data, m.id, "m", m))
    for t in tr.list_period(arc, start, end):
        items.append((t.data, t.id, "t", t))
    items.sort(key=lambda x: (x[0], x[2], x[1]))
    for _, _, kind, it in items:
        if kind == "m":
            m = it
            cur = arc.resource_currency(m.risorsa_id)["sigla"]
            out.append([m.data, name("gruppi", m.gruppo_id), name("sottogruppi", m.sottogruppo_id),
                        name("risorse", m.risorsa_id), cur, m.entrata or None, m.uscita or None,
                        "*" if m.saldato else "", name("modalita", m.modalita_id),
                        name("destinazioni", m.destinazione_id), name("soggetti", m.soggetto_id),
                        m.descrizione, m.rif, m.controllo])
        else:
            t = it
            legs = [(t.prelievo, True)] + [(v, False) for v in t.versamenti]
            for i, (leg, is_out) in enumerate(legs):
                cur = arc.resource_currency(leg.risorsa_id)["sigla"]
                out.append([t.data, TRANSFER_MARK if i == 0 else "", "", name("risorse", leg.risorsa_id), cur,
                            None if is_out else leg.importo, leg.importo if is_out else None,
                            "*" if t.saldato else "", "", "", "", t.descrizione, "", ""])
    return out


def export_xlsx(arc: Archive, path: Path | str, start: date, end: date):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimenti"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    fmt = "#,##0." + "0" * arc.decimals if arc.decimals else "#,##0"
    for r in export_rows(arc, start, end):
        ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
        row = ws.max_row
        ws.cell(row, 1).number_format = "DD/MM/YYYY"
        ws.cell(row, 6).number_format = fmt
        ws.cell(row, 7).number_format = fmt
    for col, w in zip("ABCDEFGHIJKLMN", [11, 22, 22, 18, 7, 12, 12, 3, 16, 16, 16, 40, 14, 3]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def _text_cells(arc: Archive, r: list) -> list[str]:
    out = []
    for v in r:
        if isinstance(v, date):
            out.append(v.strftime("%d/%m/%Y"))
        elif isinstance(v, Decimal):
            out.append(fmt_amount(v, arc.decimals))
        else:
            out.append("" if v is None else str(v))
    return out


def export_csv(arc: Archive, path: Path | str, start: date, end: date, delimiter: str = "\t"):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delimiter)
        w.writerow(COLUMNS)
        for r in export_rows(arc, start, end):
            w.writerow(_text_cells(arc, r))


def export_html(arc: Archive, path: Path | str, start: date, end: date, title: str = "Movimenti"):
    rows = "\n".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in _text_cells(arc, r)) + "</tr>"
                     for r in export_rows(arc, start, end))
    head = "".join(f"<th>{html.escape(c)}</th>" for c in COLUMNS)
    Path(path).write_text(
        f"<!doctype html><meta charset='utf-8'><title>{html.escape(title)}</title>"
        "<style>body{font-family:sans-serif}table{border-collapse:collapse}td,th{border:1px solid #bbb;"
        "padding:2px 6px}td:nth-child(6),td:nth-child(7){text-align:right}</style>"
        f"<h1>{html.escape(title)}</h1><p>Dal {start:%d/%m/%Y} al {end:%d/%m/%Y}</p>"
        f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>", encoding="utf-8")


# --- importazione -----------------------------------------------------------
def _norm_header(v) -> str | None:
    if v is None:
        return None
    h = str(v).strip().lower()
    h = ALIASES.get(h, h)
    return h if h in COLUMNS else None


def _to_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"data non riconosciuta: {v!r}")


def _to_amount(v) -> Decimal:
    if v is None or v == "":
        return ZERO
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    return parse_amount(str(v))


def _cell_text(v) -> str:
    return "" if v is None else str(v).strip()


def read_sheet_rows(path: Path | str) -> list[tuple]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def import_rows(arc: Archive, rows: list[tuple], report: ImportReport | None = None) -> ImportReport:
    """Importa righe di un foglio (lista di tuple) secondo le regole del modulo."""
    report = report or ImportReport()
    header_idx = None
    for i, r in enumerate(rows[:HEADER_SEARCH_ROWS]):
        if r and _norm_header(r[0]):
            header_idx = i
            break
    if header_idx is None:
        raise ValidationError("Intestazione non trovata: la prima colonna di una delle prime 20 righe deve "
                              "contenere il nome di un campo (data, gruppo, risorsa, ...)")
    header = [_norm_header(h) for h in rows[header_idx]]
    missing = {"data", "gruppo", "risorsa"} - set(header)
    if missing:
        raise ValidationError("Colonne obbligatorie mancanti: " + ", ".join(sorted(missing)))
    ignored = [str(h) for h, n in zip(rows[header_idx], header) if h not in (None, "") and n is None]
    if ignored:
        report.warn("Colonne ignorate: " + ", ".join(ignored))
    if "valuta" in header:
        report.warn("La colonna valuta è ignorata: vale la valuta della risorsa")

    res = Resolver(arc, report)
    pending_transfer: list[tuple[int, dict]] = []

    def flush_transfer():
        if not pending_transfer:
            return
        legs = pending_transfer[:]
        pending_transfer.clear()
        outs = [(n, r) for n, r in legs if r["uscita"] > 0]
        ins = [(n, r) for n, r in legs if r["entrata"] > 0]
        first = legs[0][0]
        if len(outs) != 1 or not 1 <= len(ins) <= tr.MAX_VERSAMENTI:
            report.warn(f"Riga {first}: giroconto non valido (serve 1 uscita e 1–3 entrate), non importato")
            return
        t = tr.Transfer(outs[0][1]["data"], tr.Leg(res.risorsa(outs[0][1]["risorsa"]), outs[0][1]["uscita"]),
                        [tr.Leg(res.risorsa(r["risorsa"]), r["entrata"]) for _, r in ins],
                        saldato=outs[0][1]["s"], descrizione=outs[0][1]["descrizione"])
        try:
            tr.save(arc, t, check_balance=False)
            report.add("Giroconti importati")
        except ValidationError as e:
            report.warn(f"Riga {first}: {e}")

    with arc.tx():
        for n, raw in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if raw is None or all(v is None or str(v).strip() == "" for v in raw):
                break
            rec = {h: (raw[i] if i < len(raw) else None) for i, h in enumerate(header) if h}
            try:
                r = {
                    "data": _to_date(rec.get("data")),
                    "gruppo": _cell_text(rec.get("gruppo")),
                    "sottogruppo": _cell_text(rec.get("sottogruppo")),
                    "risorsa": _cell_text(rec.get("risorsa")),
                    "entrata": abs(_to_amount(rec.get("entrata"))),
                    "uscita": abs(_to_amount(rec.get("uscita"))),
                    "s": _cell_text(rec.get("s")).lower() in ("*", "x", "s", "si", "sì", "1", "true", "vero"),
                    "modalita": _cell_text(rec.get("modalità")),
                    "destinazione": _cell_text(rec.get("destinazione")),
                    "soggetto": _cell_text(rec.get("soggetto")),
                    "descrizione": _cell_text(rec.get("descrizione"))[:mv.MAX_DESCRIZIONE],
                    "rif": _cell_text(rec.get("rif"))[:mv.MAX_RIF],
                    "c": _cell_text(rec.get("c"))[:1],
                }
            except ValueError as e:
                report.warn(f"Riga {n}: {e}, saltata")
                continue
            if r["gruppo"].lower() == TRANSFER_MARK:
                flush_transfer()
                pending_transfer.append((n, r))
                continue
            if not r["gruppo"] and pending_transfer:
                pending_transfer.append((n, r))
                continue
            flush_transfer()
            if not r["gruppo"] or not r["risorsa"]:
                report.warn(f"Riga {n}: gruppo e risorsa sono obbligatori, saltata")
                continue
            if r["entrata"] and r["uscita"]:
                report.warn(f"Riga {n}: entrata e uscita insieme, saltata")
                continue
            importo = r["entrata"] - r["uscita"]
            gid = res.gruppo(r["gruppo"], 1 if importo > 0 else -1 if importo < 0 else 0)
            rid = res.risorsa(r["risorsa"])
            m = mv.Movement(
                data=r["data"], risorsa_id=rid, gruppo_id=gid, importo=importo, saldato=r["s"],
                sottogruppo_id=res.sottogruppo(gid, r["sottogruppo"]) if r["sottogruppo"] else None,
                modalita_id=res.modalita(rid, r["modalita"]) if r["modalita"] else None,
                destinazione_id=res.simple("destinazioni", r["destinazione"], "Destinazioni create")
                if r["destinazione"] else None,
                soggetto_id=res.simple("soggetti", r["soggetto"], "Soggetti creati") if r["soggetto"] else None,
                descrizione=r["descrizione"], rif=r["rif"], controllo=r["c"])
            try:
                mv.save(arc, m)
                report.add("Movimenti importati")
            except ValidationError as e:
                report.warn(f"Riga {n}: {e}")
        flush_transfer()
    return report


def import_xlsx(arc: Archive, path: Path | str) -> ImportReport:
    return import_rows(arc, read_sheet_rows(path))


# --- esportazione di un prospetto qualsiasi ----------------------------------------------
_IT_NUMBER = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+(,\d+)?$")


def _cell_value(v):
    """Riconverte in numero i testi formattati all'italiana ("1.234,50", "12,5%")."""
    if not isinstance(v, str):
        return float(v) if isinstance(v, Decimal) else v
    t = v.strip()
    if t.endswith("%") and _IT_NUMBER.match(t[:-1].strip()):
        return float(t[:-1].strip().replace(".", "").replace(",", ".")) / 100
    if _IT_NUMBER.match(t):
        return float(t.replace(".", "").replace(",", "."))
    return v


def export_table_xlsx(path: Path | str, title: str, headers: list[str], rows: list[list], note: str = ""):
    """Scrive un prospetto (intestazioni + righe) in un foglio Excel con titolo e numeri veri."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospetto"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    if note:
        ws.append([note])
    ws.append([])
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    head_row = ws.max_row
    for r in rows:
        vals = [_cell_value(v) for v in r]
        ws.append(vals)
        for i, (raw, v) in enumerate(zip(r, vals), 1):
            if isinstance(v, float):
                pct = isinstance(raw, str) and raw.strip().endswith("%")
                ws.cell(ws.max_row, i).number_format = "0.0%" if pct else "#,##0.00"
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)]) + 2
        ws.column_dimensions[ws.cell(head_row, i).column_letter].width = min(width, 50)
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    wb.save(path)
